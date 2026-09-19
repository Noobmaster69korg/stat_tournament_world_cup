import streamlit as st
import sqlite3
import pandas as pd
import json
import requests
from bs4 import BeautifulSoup
from itertools import combinations
import os
import io
import glob
import subprocess
import re
from pypdf import PdfReader

# --- 1. DATABASE & CONNECTION MANAGER ---
# Scraped datasets are now named (one .db file per name) instead of a single
# shared 'scraped_stats.db'. Each browser tab picks its own "Active Dataset"
# in its own session_state, so two tabs can independently work on e.g. Test
# stats and ODI stats at the same time without stomping on each other.
LOCAL_DB_FILE = 'cricket_stats.db'
RESERVED_NAMES = {"Local"}

def safe_dataset_name(raw):
    """Sanitize a user-typed dataset name into a safe filename stem."""
    name = "".join(c for c in raw.strip() if c.isalnum() or c in ("_", "-"))
    name = name or "scraped_stats"
    if name in RESERVED_NAMES:
        name = f"{name}_data"
    return name

def list_scraped_datasets():
    """All named scraped datasets currently on disk (excludes the local DB)."""
    names = []
    for f in glob.glob("*.db"):
        stem = f[:-3]
        if f == LOCAL_DB_FILE or stem in RESERVED_NAMES:
            continue
        names.append(stem)
    return sorted(names)

def get_db_info():
    """Returns (connection, time_column_name) for THIS tab's active dataset selection."""
    active = st.session_state.get("active_dataset", "Local")
    if active != "Local":
        path = f"{active}.db"
        if os.path.exists(path):
            return sqlite3.connect(path, check_same_thread=False), "Year"
        # Dataset was deleted (e.g. by another tab) since this tab last checked.
        st.session_state.active_dataset = "Local"
    return sqlite3.connect(LOCAL_DB_FILE, check_same_thread=False), "Season"

# --- 1.4 GIT PERSISTENCE HELPERS ---
# Streamlit Cloud's disk is ephemeral: anything written at runtime (a built or
# uploaded .db file) is wiped whenever the app sleeps/restarts/redeploys.
# The only way for a dataset to survive that is to actually be committed into
# the GitHub repo the app deploys from (exactly like cricket_stats.db already
# is). These helpers commit/remove a dataset file via git + push, using a
# GitHub token supplied through Streamlit's secrets (Settings -> Secrets):
#
#   GITHUB_TOKEN = "ghp_xxx..."      # a token with 'repo' write access
#   GITHUB_REPO  = "username/repo"   # the repo this app deploys from
#   GITHUB_BRANCH = "main"           # optional, defaults to "main"
#
# NOTE: pushing a new commit will make Streamlit Cloud auto-redeploy the app
# shortly after (that's how Streamlit Cloud detects repo changes) - the app
# will briefly restart, which is expected.
def _git_creds():
    token = st.secrets.get("GITHUB_TOKEN") if hasattr(st, "secrets") else None
    repo = st.secrets.get("GITHUB_REPO") if hasattr(st, "secrets") else None
    branch = st.secrets.get("GITHUB_BRANCH", "main") if hasattr(st, "secrets") else "main"
    return token, repo, branch

def git_is_tracked(path):
    """True if this file is already committed in the repo (i.e. already permanent)."""
    try:
        result = subprocess.run(["git", "ls-files", "--error-unmatch", path], capture_output=True, text=True)
        return result.returncode == 0
    except Exception:
        return False

def _git_run(args):
    return subprocess.run(["git"] + args, capture_output=True, text=True)

def make_dataset_permanent(path, message):
    """Commits + pushes a dataset file to GitHub so it survives restarts forever."""
    token, repo, branch = _git_creds()
    if not token or not repo:
        return False, "Missing GITHUB_TOKEN / GITHUB_REPO in this app's Secrets (Settings -> Secrets on Streamlit Cloud)."
    try:
        _git_run(["config", "user.email", "app@streamlit.local"])
        _git_run(["config", "user.name", "Cricket Stats App"])
        _git_run(["add", path])
        commit = _git_run(["commit", "-m", message])
        if commit.returncode != 0 and "nothing to commit" not in (commit.stdout + commit.stderr):
            return False, commit.stderr or commit.stdout
        remote_url = f"https://{token}@github.com/{repo}.git"
        push = _git_run(["push", remote_url, f"HEAD:{branch}"])
        if push.returncode != 0:
            return False, push.stderr or push.stdout
        return True, "Committed and pushed. The app will redeploy shortly to reflect this."
    except Exception as e:
        return False, str(e)

def remove_dataset_permanently(path, message):
    """Removes a dataset file from the repo on GitHub (undoes make_dataset_permanent)."""
    token, repo, branch = _git_creds()
    if not token or not repo:
        return False, "Missing GITHUB_TOKEN / GITHUB_REPO in this app's Secrets (Settings -> Secrets on Streamlit Cloud)."
    try:
        _git_run(["config", "user.email", "app@streamlit.local"])
        _git_run(["config", "user.name", "Cricket Stats App"])
        _git_run(["rm", "--cached", path])
        commit = _git_run(["commit", "-m", message])
        if commit.returncode != 0 and "nothing to commit" not in (commit.stdout + commit.stderr):
            return False, commit.stderr or commit.stdout
        remote_url = f"https://{token}@github.com/{repo}.git"
        push = _git_run(["push", remote_url, f"HEAD:{branch}"])
        if push.returncode != 0:
            return False, push.stderr or push.stdout
        return True, "Removed from GitHub and pushed. The app will redeploy shortly."
    except Exception as e:
        return False, str(e)

# --- 1.5 CUSTOMIZABLE METRIC CONFIG ---
# Instead of hardcoding Runs/Ave/SR (batting) and Wkts/Ave/Econ (bowling), the
# user picks any 3 columns per discipline plus a direction for each. Every SQL
# builder below is driven purely by this config.

DEFAULT_BAT_METRICS = [
    {"col": "Runs", "direction": "higher_better"},
    {"col": "Ave", "direction": "higher_better"},
    {"col": "SR", "direction": "higher_better"},
]
DEFAULT_BOWL_METRICS = [
    {"col": "Wkts", "direction": "higher_better"},
    {"col": "Ave", "direction": "lower_better"},
    {"col": "Econ", "direction": "lower_better"},
]
LOWER_BETTER_HINTS = ("ave", "econ", "sr", "rpo", "average", "economy")

def guess_direction(col, discipline):
    """Best-effort default so the UI doesn't start every column as 'higher is
    better'. The user can always override via the radio button."""
    if discipline == "bowling" and any(k in col.lower() for k in LOWER_BETTER_HINTS):
        return "lower_better"
    return "higher_better"

def get_table_columns(conn, table):
    """Numeric-ish columns available in a table, excluding identifiers."""
    try:
        rows = conn.execute(f"PRAGMA table_info({table})").fetchall()
    except Exception:
        return []
    cols = []
    for r in rows:
        name, decl_type = r[1], (r[2] or "").upper()
        if name in ("Player", "Season", "Year"):
            continue
        if any(t in decl_type for t in ("CHAR", "TEXT", "CLOB")):
            continue
        cols.append(name)
    return cols

def q_col(col):
    """Bracket-quote a column name so odd identifiers (e.g. '4s', '100') are safe in SQL."""
    return f"[{col}]"

def metric_picker(prefix, avail_cols, defaults, discipline_for_guess):
    """Renders 3 (column, direction) pickers and returns the chosen metric config."""
    if len(avail_cols) < 3:
        st.error(f"'{discipline_for_guess}' table needs at least 3 numeric columns to configure metrics. Found: {avail_cols}")
        return defaults
    metrics, chosen_so_far = [], []
    cols_ui = st.columns(3)
    for i in range(3):
        with cols_ui[i]:
            default_col = defaults[i]['col'] if i < len(defaults) and defaults[i]['col'] in avail_cols else avail_cols[i % len(avail_cols)]
            idx = avail_cols.index(default_col) if default_col in avail_cols else 0
            col_choice = st.selectbox(f"Metric {i+1}", avail_cols, index=idx, key=f"{prefix}_col_{i}")
            default_dir = next((d['direction'] for d in defaults if d['col'] == col_choice), None)
            if default_dir is None:
                default_dir = guess_direction(col_choice, discipline_for_guess)
            dir_choice = st.radio(
                "Direction", ["Higher is better", "Lower is better"],
                index=0 if default_dir == "higher_better" else 1,
                key=f"{prefix}_dir_{i}", horizontal=True,
            )
            metrics.append({"col": col_choice, "direction": "higher_better" if dir_choice == "Higher is better" else "lower_better"})
            chosen_so_far.append(col_choice)
    if len(set(chosen_so_far)) < 3:
        st.warning("You've selected the same column more than once — comparisons involving it will be redundant.")
    return metrics

def build_case_sql(metrics, thresholds, op):
    """Generic Wins/Ties/Losses/Meets SQL fragment vs. fixed threshold values.
    op: 'win' (strictly better), 'tie' (exactly equal), 'loss' (strictly worse),
        'meets' (>=/<= threshold, used for 'consistency' style checks)."""
    parts = []
    for m, t in zip(metrics, thresholds):
        col, higher = m['col'], m['direction'] == 'higher_better'
        if op == 'tie':
            cmp = '='
        elif op == 'win':
            cmp = '>' if higher else '<'
        elif op == 'loss':
            cmp = '<' if higher else '>'
        elif op == 'meets':
            cmp = '>=' if higher else '<='
        else:
            raise ValueError(f"Unknown op: {op}")
        parts.append(f"(CASE WHEN {q_col(col)} {cmp} {t} THEN 1 ELSE 0 END)")
    return " + ".join(parts)

def build_pairwise_sql(metrics, alias_a, alias_b, op):
    """Generic row-vs-row (e.g. A.Runs > B.Runs) SQL fragment. op: 'win' or 'loss'."""
    parts = []
    for m in metrics:
        col, higher = m['col'], m['direction'] == 'higher_better'
        cmp = ('>' if higher else '<') if op == 'win' else ('<' if higher else '>')
        parts.append(f"(CASE WHEN {alias_a}.{q_col(col)} {cmp} {alias_b}.{q_col(col)} THEN 1 ELSE 0 END)")
    return "(" + " + ".join(parts) + ")"

def build_omit_sql(t_col, omit_set, table, alias=None):
    """WHERE-clause fragment excluding specific (table, Player, Year) triples
    from a squad's eligible pool - e.g. a specific bad year the user chose to
    leave out without dropping the player from the squad entirely. Only
    triples matching `table` (batting/bowling) apply; alias is the SQL table
    alias to qualify Player/{t_col} with (None for an unaliased query)."""
    prefix = f"{alias}." if alias else ""
    conds = []
    for tbl, player, year in omit_set:
        if tbl != table:
            continue
        p = str(player).replace("'", "''")
        y = str(year).replace("'", "''")
        conds.append(f"({prefix}Player = '{p}' AND {prefix}{t_col} = '{y}')")
    if not conds:
        return ""
    return " AND NOT (" + " OR ".join(conds) + ")"

def compare_rows(row_a, row_b, metrics):
    """Python-side equivalent of build_pairwise_sql, for the Format Analysis loops."""
    wins = 0
    for m in metrics:
        col, higher = m['col'], m['direction'] == 'higher_better'
        va, vb = row_a[col], row_b[col]
        if (higher and va > vb) or ((not higher) and va < vb):
            wins += 1
    return wins

# --- 1.6 YEAR / COUNTRY FILTERING ---
# Neither table has a dedicated Country column - it's baked into the player
# name as a trailing code (e.g. "SL Malinga SL", "Umar Gul PAK"). We detect it
# heuristically and only surface the country filter when the dataset actually
# looks like it follows that convention.
def extract_year_num(val):
    """Pulls a 4-digit year out of a Season/Year value (handles '2019', '2019/20', etc.)."""
    m = re.search(r"(\d{4})", str(val))
    return int(m.group(1)) if m else None

def extract_country(player_name):
    """Best-effort: the last whitespace-separated token of the player name."""
    if not isinstance(player_name, str) or not player_name.strip():
        return "Unknown"
    tokens = player_name.strip().split()
    return tokens[-1].upper() if tokens else "Unknown"

def country_filter_is_meaningful(all_players, max_distinct=40):
    """If almost every player has a 'unique country', it's not really a country
    code - it's just their surname. Only offer the filter when the number of
    distinct suffixes is small relative to typical country-code counts."""
    if not all_players:
        return False, []
    codes = sorted(set(extract_country(p) for p in all_players))
    return (len(codes) <= max_distinct), codes

def build_filtered_connection(raw_conn, t_col, year_range, countries):
    """Reads batting/bowling from raw_conn, applies year/country filters, and
    returns a NEW in-memory connection with tables of the SAME names
    ('batting'/'bowling') - so every existing query elsewhere in the app keeps
    working unchanged against this filtered view."""
    bat_df = pd.read_sql("SELECT * FROM batting", raw_conn)
    bowl_df = pd.read_sql("SELECT * FROM bowling", raw_conn)

    for df in (bat_df, bowl_df):
        df['__year_num'] = df[t_col].apply(extract_year_num)
        df['__country'] = df['Player'].apply(extract_country)

    y_min, y_max = year_range
    bat_df = bat_df[bat_df['__year_num'].notna() & bat_df['__year_num'].between(y_min, y_max)]
    bowl_df = bowl_df[bowl_df['__year_num'].notna() & bowl_df['__year_num'].between(y_min, y_max)]

    if countries:
        bat_df = bat_df[bat_df['__country'].isin(countries)]
        bowl_df = bowl_df[bowl_df['__country'].isin(countries)]

    bat_df = bat_df.drop(columns=['__year_num', '__country'])
    bowl_df = bowl_df.drop(columns=['__year_num', '__country'])

    mem_conn = sqlite3.connect(":memory:", check_same_thread=False)
    bat_df.to_sql("batting", mem_conn, index=False)
    bowl_df.to_sql("bowling", mem_conn, index=False)
    return mem_conn

# --- 1.7 AUCTION SHEET HELPERS ---
# Takes an auction-order PDF (numbered list of player names) plus one or more
# format's batting/bowling "Global Rankings" CSVs (Player, Year, Wins %,
# Losses, Ties), and produces one row per auction player with loss-threshold
# counts and a best-year rank for every format/discipline combination included.
_PAGE_FOOTER_RE = re.compile(r"^page\s+\d+\s+of\s+\d+$", re.IGNORECASE)

def parse_auction_order_bytes(file_bytes, filename):
    """Parses an uploaded auction order file into an ordered list of
    (order, player_name) tuples. Accepts either a .txt file (read as plain
    text - most reliable) or a .pdf (text extracted via pypdf).

    PDF text extraction is NOT always reliable: some PDFs (e.g. exported from
    certain apps, or with embedded/subsetted fonts) don't yield clean,
    extractable body text via pypdf even though the text is visibly there and
    selectable - sometimes the only text that comes out is incidental stuff
    like an auto-inserted page-number footer. If that happens here, uploading
    a plain .txt file with the same player list (one per line) sidesteps PDF
    parsing entirely and is much more dependable.
    """
    if filename.lower().endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="replace")
    else:
        reader = PdfReader(io.BytesIO(file_bytes))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)

    players = []
    seq = 0
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if _PAGE_FOOTER_RE.match(line):
            # Common PDF footer artifact ("Page 2 of 3") - never a real player.
            continue
        m = re.match(r"^(\d+)\s+(.+?)\s*$", line)
        if m:
            players.append((int(m.group(1)), m.group(2).strip()))
        else:
            # No leading order number on this line (e.g. a plain unnumbered
            # name list) - fall back to treating it as a name and assigning
            # order based on its position in the document, instead of
            # silently dropping it.
            seq += 1
            players.append((seq, line))
    players.sort(key=lambda x: x[0])
    return players

# Name matching: CSVs carry a trailing country code (e.g. "SL Malinga SL"),
# auction names don't. We match on surname + given-name initials, with an
# explicit confidence tier - anything not a unique, full-initials match gets
# flagged for manual review rather than silently trusted (a first-initial-only
# match once produced a genuine false positive: "Shaheen Shah Afridi" ->
# "Shahid Afridi" - so that tier is always labeled, never called "exact").
_NAME_PREFIXES = {"de", "van", "der", "al", "bin", "du", "le"}

def _split_name(name):
    tokens = [t for t in re.sub(r"[^\w\s'\-]", "", name).split() if t]
    if not tokens:
        return None
    surname_tokens = [tokens[-1]]
    i = len(tokens) - 2
    while i >= 0 and tokens[i].lower() in _NAME_PREFIXES:
        surname_tokens.insert(0, tokens[i])
        i -= 1
    surname = " ".join(surname_tokens).lower()
    given = tokens[: len(tokens) - len(surname_tokens)]
    initials = "".join(g[0].lower() for g in given)
    return surname, initials

def _build_name_index(player_series):
    by_full, by_surname_firstinit, by_surname = {}, {}, {}
    for full in player_series.dropna().unique():
        tokens = str(full).split()
        core = " ".join(tokens[:-1]) if len(tokens) > 1 else full  # strip trailing country code
        parsed = _split_name(core)
        if parsed is None:
            continue
        surname, initials = parsed
        by_full.setdefault((surname, initials), []).append(full)
        if initials:
            by_surname_firstinit.setdefault((surname, initials[0]), []).append(full)
        by_surname.setdefault(surname, []).append(full)
    return by_full, by_surname_firstinit, by_surname

def _match_player(auction_name, idx):
    by_full, by_sfi, by_s = idx
    parsed = _split_name(auction_name)
    if parsed is None:
        return None, "no_match"
    surname, initials = parsed
    cands = by_full.get((surname, initials), [])
    if len(cands) == 1: return cands[0], "exact"
    if len(cands) > 1: return cands[0], "ambiguous"
    cands2 = by_sfi.get((surname, initials[0] if initials else ""), [])
    if len(cands2) == 1: return cands2[0], "weak"
    if len(cands2) > 1: return cands2[0], "ambiguous"
    cands3 = by_s.get(surname, [])
    if len(cands3) == 1: return cands3[0], "weak"
    if len(cands3) > 1: return cands3[0], "ambiguous"
    return None, "no_match"

def _parse_raw_loss(loss_str):
    m = re.match(r"\s*(\d+)", str(loss_str))
    return int(m.group(1)) if m else None

def _player_loss_stats(matched_name, df, thresholds):
    rows = df[df['Player'] == matched_name].copy()
    if rows.empty:
        return None
    rows['_raw_loss'] = rows['Losses'].apply(_parse_raw_loss)
    rows = rows.dropna(subset=['_raw_loss'])
    if rows.empty:
        return None
    out = {f"<{th}": int((rows['_raw_loss'] < th).sum()) for th in thresholds}
    best_idx = rows['_raw_loss'].idxmin()
    best_row = rows.loc[best_idx]
    out['best_year'] = best_row['Year']
    out['best_year_losses'] = int(best_row['_raw_loss'])
    all_raw = df['Losses'].apply(_parse_raw_loss).dropna()
    out['rank'] = int((all_raw < out['best_year_losses']).sum()) + 1  # competition ranking - lower losses = better rank
    out['total_n'] = int(all_raw.shape[0])
    return out

AUCTION_KEEP_AUTO = "↩️ Keep automatic match"
AUCTION_EXCLUDE = "❌ Exclude (no match)"

def _load_format_dataframes(formats_config):
    """Reads each configured format's batting/bowling CSVs into DataFrames,
    keyed by (format_name, 'bat'/'bowl'). Cached in session_state after the
    initial build so manual corrections below don't depend on the uploaded
    file widgets still holding their content."""
    loaded = {}
    for fmt_name, cfg in formats_config.items():
        for disc in ("bat", "bowl"):
            src = cfg[disc]
            src.seek(0)
            df = pd.read_csv(src)
            df.columns = [c.strip() for c in df.columns]
            loaded[(fmt_name, disc)] = df
    return loaded

def _build_indices(loaded):
    """One name-index per (format, discipline) sheet, built once and reused
    both for the initial automatic match and for cross-sheet propagation."""
    return {key: _build_name_index(df['Player']) for key, df in loaded.items()}

def _auto_match_records(auction_players, indices):
    """One independent automatic name-match attempt per (auction player,
    format, discipline) cell - the starting point before cross-sheet
    propagation or any manual corrections."""
    records = {}
    for order, name in auction_players:
        for fmt_name, disc in indices:
            matched, status = _match_player(name, indices[(fmt_name, disc)])
            records[(order, fmt_name, disc)] = {
                "order": order, "auction_player": name,
                "format": fmt_name, "disc": disc,
                "matched_name": matched, "status": status,
            }
    return records

def _core_surname(matched_name):
    """Surname of a matched CSV player string, ignoring the trailing country code."""
    tokens = str(matched_name).split()
    core = " ".join(tokens[:-1]) if len(tokens) > 1 else matched_name
    parsed = _split_name(core)
    return parsed[0] if parsed else None

def _resolve_all(auction_players, format_names, indices, auto_records, overrides):
    """Resolves every (auction player, format, discipline) cell to its
    effective match, then propagates identity across sheets for the same
    player: once a player is confidently identified (an exact automatic
    match, or a manual correction) in ANY one format/discipline, that same
    surname is used to resolve them in every OTHER sheet too - rather than
    treating each sheet's match as fully independent. This matters because
    the same real player can come out as an 'exact' match in one sheet but
    only 'weak'/'ambiguous' in another purely due to how that sheet's CSV
    happened to write their name (e.g. one scrape uses full initials, another
    a single initial) - country suffix differences are already ignored by
    the surname/initials matching itself, this handles initials-formatting
    differences across separately-scraped sheets.

    A cell only gets upgraded this way if the confirmed surname is UNIQUE
    within that other sheet too - if two different players share that surname
    there, we can't safely disambiguate and it's left for manual review.
    Manual overrides (including an explicit 'exclude') are never touched by
    propagation.

    Returns (resolved, confirmed_counts):
      resolved: {(order, fmt, disc): (matched_name_or_None, status)}
        status one of: 'exact', 'manual', 'manually_excluded',
        'cross_confirmed', 'weak', 'ambiguous', 'no_match'
      confirmed_counts: {order: count of cells resolved with real confidence}
    """
    resolved = {}
    for order, name in auction_players:
        for fmt_name in format_names:
            for disc in ("bat", "bowl"):
                key = (order, fmt_name, disc)
                if key in overrides:
                    chosen = overrides[key]
                    resolved[key] = (None, "manually_excluded") if chosen == AUCTION_EXCLUDE else (chosen, "manual")
                else:
                    rec = auto_records[key]
                    resolved[key] = (rec["matched_name"], rec["status"])

    # Cross-sheet propagation.
    for order, name in auction_players:
        confirmed_surnames = []
        for fmt_name in format_names:
            for disc in ("bat", "bowl"):
                matched, status = resolved[(order, fmt_name, disc)]
                if status in ("exact", "manual") and matched:
                    surname = _core_surname(matched)
                    if surname:
                        confirmed_surnames.append(surname)
        if not confirmed_surnames:
            continue
        surname = max(set(confirmed_surnames), key=confirmed_surnames.count)
        for fmt_name in format_names:
            for disc in ("bat", "bowl"):
                key = (order, fmt_name, disc)
                if key in overrides:
                    continue  # never override an explicit user choice
                matched, status = resolved[key]
                if status in ("exact", "manual"):
                    continue  # already confident on its own terms
                by_surname = indices[(fmt_name, disc)][2]
                cands = by_surname.get(surname, [])
                if len(cands) == 1:
                    resolved[key] = (cands[0], "cross_confirmed")

    confirmed_counts = {}
    for order, name in auction_players:
        confirmed_counts[order] = sum(
            1 for fmt_name in format_names for disc in ("bat", "bowl")
            if resolved[(order, fmt_name, disc)][1] in ("exact", "manual", "cross_confirmed")
        )
    return resolved, confirmed_counts

def _needs_review(status, order, confirmed_counts):
    """Whether a cell's match is worth flagging for a human to check.
    - Ambiguous/weak matches always are: there IS a row there, just an
      uncertain one, and cross-sheet propagation (see _resolve_all) couldn't
      resolve it either.
    - A 'no match found' is only worth flagging if this player wasn't
      identified in ANY format/discipline at all. If they were found
      elsewhere, a no-match here almost always just means they have no
      eligible years/appearances in that particular format or discipline
      (e.g. a specialist batter absent from the bowling sheet) - not a
      real data problem.
    - 'cross_confirmed' (and 'exact'/'manual') never need review - that's
      the whole point: a confident match anywhere resolves it everywhere."""
    if status in ("weak", "ambiguous"):
        return True
    if status == "no_match":
        return confirmed_counts.get(order, 0) == 0
    return False

def build_auction_sheet_df(auction_players, format_names, loaded, resolved, confirmed_counts, thresholds):
    """Builds the result DataFrame + a list of human-readable warning strings
    from an already-resolved match set (see _resolve_all)."""
    rows, warnings = [], []
    for order, name in auction_players:
        row = {"Auction Order": order, "Auction Player": name}
        notes = []
        for fmt_name in format_names:
            for disc, disc_label in (("bat", "Bat"), ("bowl", "Bowl")):
                df = loaded[(fmt_name, disc)]
                prefix = f"{fmt_name} {disc_label}"
                matched, status = resolved[(order, fmt_name, disc)]
                stats = _player_loss_stats(matched, df, thresholds) if matched else None
                if stats is None:
                    for th in thresholds:
                        row[f"{prefix} <{th}"] = None
                    row[f"{prefix} Best Year"] = None
                    row[f"{prefix} Best Year Losses"] = None
                    row[f"{prefix} Rank"] = None
                else:
                    for th in thresholds:
                        row[f"{prefix} <{th}"] = stats[f"<{th}"]
                    row[f"{prefix} Best Year"] = stats['best_year']
                    row[f"{prefix} Best Year Losses"] = stats['best_year_losses']
                    row[f"{prefix} Rank"] = f"{stats['rank']} of {stats['total_n']}"
                if status == "manual":
                    notes.append(f"{prefix}: manually corrected to '{matched}'")
                elif status == "manually_excluded":
                    notes.append(f"{prefix}: manually excluded (no match)")
                elif status == "cross_confirmed":
                    notes.append(f"{prefix}: matched '{matched}' (confirmed via another format)")
                elif status == "no_match":
                    if _needs_review(status, order, confirmed_counts):
                        notes.append(f"{prefix}: no match found, please verify")
                    # else: found in another format/discipline already - most likely
                    # this player simply has no eligible years here, not worth a note.
                elif status != "exact":
                    notes.append(f"{prefix}: matched '{matched}' ({status}, please verify)")
        row["Match Notes"] = "; ".join(notes)
        if notes:
            warnings.append(f"#{order} {name}: {row['Match Notes']}")
        rows.append(row)
    return pd.DataFrame(rows), warnings



# --- 2. ROBUST PAGINATED SCRAPER (Always creates 'Year' column) ---
def scrape_full_cricinfo(base_url, discipline):
    all_data = []
    page = 1
    progress_container = st.sidebar.empty()
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'}

    while True:
        url = f"{base_url};page={page}"
        try:
            response = requests.get(url, headers=headers, timeout=15)
            if response.status_code != 200: break
            tables = pd.read_html(io.StringIO(response.text))
            df_page = None
            for t in tables:
                if 'Player' in t.columns:
                    df_page = t
                    break
            if df_page is None or df_page.empty: break

            if 'Span' in df_page.columns: df_page = df_page.rename(columns={'Span': 'Year'})
            if 'Player' not in df_page.columns or 'Year' not in df_page.columns: break

            # Keep every usable column from the page (not just a fixed subset)
            # so any of them can later be picked as a comparison metric.
            drop_like = [c for c in df_page.columns if str(c).lower().startswith('unnamed')]
            df_final = df_page.drop(columns=drop_like, errors='ignore').copy()
            df_final['Player'] = df_final['Player'].str.replace(r'[^\w\s]', '', regex=True).str.strip()
            all_data.append(df_final)
            progress_container.info(f"Scraped {discipline} Page {page}...")
            if len(df_page) < 50: break
            page += 1
            if page > 40: break
        except Exception:
            break

    if not all_data: return None
    full_df = pd.concat(all_data, ignore_index=True)
    for col in full_df.columns:
        if col in ("Player", "Year"): continue
        full_df[col] = pd.to_numeric(full_df[col], errors='coerce').fillna(0)
    return full_df

# --- 3. GLOBAL HELPERS ---
st.set_page_config(page_title="Cricket Stats Engine", layout="wide")

def fmt(count, total):
    if total <= 0: return "0 (0.0%)"
    perc = (count * 100.0 / total)
    return f"{int(count)} ({perc:.1f}%)"

def pct(count, total):
    """Like fmt(), but returns a plain float - use this for any column that
    goes into an interactive st.dataframe, so clicking the header to sort
    actually sorts numerically instead of alphabetically on a text string."""
    if total <= 0: return 0.0
    return round(count * 100.0 / total, 1)

def get_profile_label(w, t, l):
    if w == 3: return "🏆 Beat all 3 categories"
    if w == 2 and t == 1: return "⭐ Beat 2 categories, Tied 1 category"
    if w == 2 and l == 1: return "✅ Beat 2 categories, Lost 1 category"
    return "Other combinations"

def display_styled_results(df, title_prefix):
    if df is None or df.empty:
        st.warning("No players found matching these criteria.")
        return
    df['Result_Profile'] = df.apply(lambda r: get_profile_label(r['WinsA'], r['TiesA'], r['LossesA']), axis=1)
    st.subheader(f"📊 Summary ({title_prefix})")
    for profile, count in df['Result_Profile'].value_counts().items():
        st.write(f"- **{count}** records: {profile}")
    st.divider()
    st.subheader("📋 Respective Lists")
    for profile in sorted(df['Result_Profile'].unique(), reverse=True):
        cols_to_drop = ['WinsA', 'TiesA', 'LossesA', 'Result_Profile']
        if 'WinsB' in df.columns: cols_to_drop.append('WinsB')
        if 'TiesB' in df.columns: cols_to_drop.append('TiesB')
        st.markdown(f"#### {profile}")
        st.dataframe(df[df['Result_Profile'] == profile].drop(columns=cols_to_drop).reset_index(drop=True), use_container_width=True, hide_index=True)

# --- 4. SIDEBAR CONFIG ---
st.sidebar.title("🌐 Live Data Bridge")
if "active_dataset" not in st.session_state: st.session_state.active_dataset = "Local"

dataset_options = ["Local"] + list_scraped_datasets()
if st.session_state.active_dataset not in dataset_options:
    st.session_state.active_dataset = "Local"  # dataset vanished (deleted elsewhere) - fall back safely

st.session_state.active_dataset = st.sidebar.selectbox(
    "📂 Active Dataset (this tab only)",
    dataset_options,
    index=dataset_options.index(st.session_state.active_dataset),
)

if st.session_state.active_dataset == "Local":
    st.sidebar.info("🏠 Active: Local Database (uses 'Season')")
else:
    st.sidebar.success(f"🟢 Active: '{st.session_state.active_dataset}' (uses 'Year')")

st.sidebar.divider()
st.sidebar.caption("Build a new dataset")
b_link = st.sidebar.text_input("ESPN Batting Link")
w_link = st.sidebar.text_input("ESPN Bowling Link")
dataset_name_input = st.sidebar.text_input("Dataset Name", placeholder="e.g. Test_2024, ODI_AUS_series")

if st.sidebar.button("🚀 Build DB"):
    if b_link and w_link:
        safe_name = safe_dataset_name(dataset_name_input)
        with st.spinner(f"Building '{safe_name}'..."):
            b_df = scrape_full_cricinfo(b_link, "batting")
            w_df = scrape_full_cricinfo(w_link, "bowling")
            if b_df is not None and w_df is not None:
                tmp_conn = sqlite3.connect(f"{safe_name}.db")
                b_df.to_sql('batting', tmp_conn, index=False, if_exists='replace')
                w_df.to_sql('bowling', tmp_conn, index=False, if_exists='replace')
                tmp_conn.close()
                st.session_state.active_dataset = safe_name  # this tab switches straight to it
                st.rerun()
            else:
                st.sidebar.error("Scraping failed — check the links and try again.")
    else:
        st.sidebar.warning("Paste both links before building.")

st.sidebar.divider()
active_db_path = LOCAL_DB_FILE if st.session_state.active_dataset == "Local" else f"{st.session_state.active_dataset}.db"
active_is_tracked = st.session_state.active_dataset == "Local" or (os.path.exists(active_db_path) and git_is_tracked(active_db_path))

st.sidebar.caption("📌 Make the active dataset permanent (commits it to GitHub — survives sleeps/restarts forever, until deleted)")
if st.session_state.active_dataset == "Local":
    st.sidebar.caption("Local database is already permanent — it's already in the repo.")
elif not os.path.exists(active_db_path):
    st.sidebar.caption("Nothing built for the active dataset yet.")
elif active_is_tracked:
    st.sidebar.success(f"'{st.session_state.active_dataset}' is already permanent.")
else:
    if st.sidebar.button(f"📌 Make '{st.session_state.active_dataset}' Permanent"):
        ok, msg = make_dataset_permanent(active_db_path, f"Add dataset: {st.session_state.active_dataset}")
        (st.sidebar.success if ok else st.sidebar.error)(msg)

st.sidebar.divider()
st.sidebar.caption("Delete the currently active dataset")
confirm_delete = st.sidebar.checkbox("Confirm delete", key="confirm_delete_active_dataset")
also_remove_github = False
if st.session_state.active_dataset != "Local" and active_is_tracked:
    also_remove_github = st.sidebar.checkbox("Also remove permanently from GitHub (can't be undone)", key="also_remove_github")
if st.sidebar.button("🗑️ Delete Active Dataset"):
    active = st.session_state.active_dataset
    if active == "Local":
        st.sidebar.warning("The Local database can't be deleted here.")
    elif not confirm_delete:
        st.sidebar.warning("Tick 'Confirm delete' first — this removes the file for every tab using it.")
    else:
        path = f"{active}.db"
        if also_remove_github:
            ok, msg = remove_dataset_permanently(path, f"Remove dataset: {active}")
            if not ok:
                st.sidebar.error(msg)
        if os.path.exists(path): os.remove(path)
        st.session_state.active_dataset = "Local"
        st.rerun()

st.sidebar.divider()
st.sidebar.caption("💾 Download a local copy (works with or without GitHub secrets configured)")
if os.path.exists(active_db_path):
    with open(active_db_path, "rb") as _f:
        st.sidebar.download_button(
            f"⬇️ Download '{st.session_state.active_dataset}'",
            data=_f.read(),
            file_name=os.path.basename(active_db_path),
            mime="application/octet-stream",
        )
else:
    st.sidebar.caption("Nothing to download for the active dataset yet.")

st.sidebar.divider()
st.sidebar.caption("⬆️ Restore a dataset you downloaded earlier")
uploaded_db = st.sidebar.file_uploader("Upload a .db file", type=["db"], key="dataset_uploader")
if uploaded_db is not None:
    default_upload_name = os.path.splitext(uploaded_db.name)[0]
    upload_name_input = st.sidebar.text_input("Save as (dataset name)", value=default_upload_name, key="upload_name_input")
    if st.sidebar.button("💾 Save Uploaded Dataset"):
        safe_name = safe_dataset_name(upload_name_input)
        with open(f"{safe_name}.db", "wb") as _f:
            _f.write(uploaded_db.getbuffer())
        st.session_state.active_dataset = safe_name
        st.rerun()

# --- 5. AUTH ---
st.title("🏏 Player Stats & Analytics Engine")
password = st.text_input("Enter Password", type="password")
if password != "qcc_stat_tourno":
    st.error("Access Denied.")
    st.stop()

# --- 6. NAV & STATE ---
if "nav_choice" not in st.session_state: st.session_state.nav_choice = "Batting Milestones"
if "squad_a" not in st.session_state: st.session_state.squad_a, st.session_state.squad_b = [], []
if "squad_a_omit" not in st.session_state: st.session_state.squad_a_omit = set()
if "squad_b_omit" not in st.session_state: st.session_state.squad_b_omit = set()
if "bat_metrics" not in st.session_state: st.session_state.bat_metrics = DEFAULT_BAT_METRICS.copy()
if "bowl_metrics" not in st.session_state: st.session_state.bowl_metrics = DEFAULT_BOWL_METRICS.copy()

conn, t_col = get_db_info()  # t_col is either 'Season' or 'Year'
raw_conn = conn  # unfiltered handle - metric config & filter-option detection read from this

with st.expander("⚙️ Configure Comparison Metrics (3 fields per discipline)", expanded=False):
    st.caption("Pick any 3 numeric fields from each table as the comparison criteria, and whether higher or lower is better for each. Everything below (Milestones, Analytics, Squad tools, Format Analysis) uses this config.")
    st.markdown("**Batting**")
    avail_bat_cols = get_table_columns(raw_conn, "batting")
    st.session_state.bat_metrics = metric_picker("bat", avail_bat_cols, st.session_state.bat_metrics, "batting")
    st.markdown("**Bowling**")
    avail_bowl_cols = get_table_columns(raw_conn, "bowling")
    st.session_state.bowl_metrics = metric_picker("bowl", avail_bowl_cols, st.session_state.bowl_metrics, "bowling")

with st.expander("🎯 Filter This Dataset (Year Range & Countries)", expanded=False):
    _all_players_df = pd.read_sql("SELECT Player FROM batting UNION SELECT Player FROM bowling", raw_conn)
    _all_years_bat = pd.read_sql(f"SELECT {t_col} as yr FROM batting", raw_conn)['yr']
    _all_years_bowl = pd.read_sql(f"SELECT {t_col} as yr FROM bowling", raw_conn)['yr']
    _year_nums = pd.concat([_all_years_bat, _all_years_bowl]).apply(extract_year_num).dropna()

    if _year_nums.empty:
        st.caption("Couldn't detect year values in this dataset — year filter unavailable.")
        year_range = (0, 9999)
    else:
        y_lo, y_hi = int(_year_nums.min()), int(_year_nums.max())
        if y_lo == y_hi:
            st.caption(f"This dataset only spans {y_lo} — nothing to range-filter.")
            year_range = (y_lo, y_hi)
        else:
            year_range = st.slider(
                "Year range", min_value=y_lo, max_value=y_hi, value=(y_lo, y_hi),
                key=f"year_range_{st.session_state.active_dataset}",
            )

    show_country_filter, detected_countries = country_filter_is_meaningful(_all_players_df['Player'].tolist())
    if show_country_filter:
        selected_countries = st.multiselect(
            "Countries (leave empty = all)", detected_countries, default=[],
            key=f"country_filter_{st.session_state.active_dataset}",
        )
    else:
        st.caption("Player names in this dataset don't look like they carry a country code, so country filtering isn't available here.")
        selected_countries = []

    filters_active = (_year_nums.empty is False and year_range != (int(_year_nums.min()), int(_year_nums.max()))) or bool(selected_countries)
    if filters_active:
        st.caption(f"✅ Filter active — {year_range[0]}–{year_range[1]}" + (f", countries: {', '.join(selected_countries)}" if selected_countries else ", all countries"))

if not _year_nums.empty:
    conn = build_filtered_connection(raw_conn, t_col, year_range, selected_countries)

nav_options = ["Batting Milestones", "Bowling Milestones", "📈 Player Analytics", "👤 Player Details", "🏟️ Squad Comparison", "🧬 Format Analysis", "✏️ Edit Data", "🏆 Auction Sheet"]
st.session_state.nav_choice = st.radio("Navigate:", nav_options, index=nav_options.index(st.session_state.nav_choice), horizontal=True)
st.divider()

# --- MILESTONES (shared by Batting & Bowling tabs) ---
def render_milestones_tab(conn, t_col, table, metrics):
    f_mode = st.radio("Display Mode:", ["Meet Set A Only", "Meet BOTH Set A and Set B"], horizontal=True, key=f"{table}_fmode")
    thresholds_key = f"{table}_thresholds"
    metric_cols = {m['col'] for m in metrics}
    if thresholds_key not in st.session_state or set(st.session_state[thresholds_key].keys()) != metric_cols:
        st.session_state[thresholds_key] = {m['col']: 0.0 for m in metrics}

    c1, c2 = st.columns(2)
    thresholds_a, thresholds_b = [], []
    with c1:
        st.caption("Set A")
        for m in metrics:
            label = f"{'Min' if m['direction']=='higher_better' else 'Max'} {m['col']} (A)"
            val = st.number_input(label, value=float(st.session_state[thresholds_key].get(m['col'], 0.0)), key=f"{table}_A_{m['col']}")
            st.session_state[thresholds_key][m['col']] = val
            thresholds_a.append(val)
    with c2:
        st.caption("Set B")
        for m in metrics:
            label = f"{'Min' if m['direction']=='higher_better' else 'Max'} {m['col']} (B)"
            val = st.number_input(label, value=0.0, key=f"{table}_B_{m['col']}")
            thresholds_b.append(val)

    select_cols = ", ".join(q_col(m['col']) for m in metrics)
    wins_a = build_case_sql(metrics, thresholds_a, 'win')
    ties_a = build_case_sql(metrics, thresholds_a, 'tie')
    losses_a = build_case_sql(metrics, thresholds_a, 'loss')
    wins_b = build_case_sql(metrics, thresholds_b, 'win')
    ties_b = build_case_sql(metrics, thresholds_b, 'tie')
    first = metrics[0]
    order_dir = "DESC" if first['direction'] == 'higher_better' else "ASC"
    q = f"""
        WITH Base AS (
            SELECT Player, {t_col} as Year, {select_cols},
                   ({wins_a}) as WinsA, ({ties_a}) as TiesA, ({losses_a}) as LossesA,
                   ({wins_b}) as WinsB, ({ties_b}) as TiesB
            FROM {table}
        )
        SELECT * FROM Base
        WHERE (WinsA + TiesA) >= 2 {"AND (WinsB + TiesB) >= 2" if "BOTH" in f_mode else ""}
        ORDER BY WinsA DESC, {q_col(first['col'])} {order_dir}
    """
    display_styled_results(pd.read_sql(q, conn), table.capitalize())

# --- TAB 1 & 2: MILESTONES ---
if st.session_state.nav_choice == "Batting Milestones":
    render_milestones_tab(conn, t_col, "batting", st.session_state.bat_metrics)

elif st.session_state.nav_choice == "Bowling Milestones":
    render_milestones_tab(conn, t_col, "bowling", st.session_state.bowl_metrics)

# --- TAB 3: ANALYTICS ---
elif st.session_state.nav_choice == "📈 Player Analytics":
    choice = st.radio("Type:", ["Career Consistency", "Global Rankings"], horizontal=True)
    disc = st.radio("Disc:", ["Batting", "Bowling"], horizontal=True).lower()
    metrics = st.session_state.bat_metrics if disc == "batting" else st.session_state.bowl_metrics
    if "Consistency" in choice:
        cols_ui = st.columns(3)
        thresholds = []
        for i, m in enumerate(metrics):
            label = f"{'Min' if m['direction']=='higher_better' else 'Max'} {m['col']}"
            val = cols_ui[i].number_input(label, value=0.0, key=f"cons_{disc}_{m['col']}")
            thresholds.append(val)
        meets_sql = build_case_sql(metrics, thresholds, 'meets')
        q = f"SELECT Player, COUNT(*) as Total, SUM(CASE WHEN ({meets_sql}) >= 2 THEN 1 ELSE 0 END) as Successful FROM {disc} GROUP BY Player HAVING Successful > 0"
        df = pd.read_sql(q, conn)
        df['Win %'] = (df['Successful'] * 100 / df['Total']).round(2)
        st.dataframe(df.sort_values("Win %", ascending=False).reset_index(drop=True), use_container_width=True, hide_index=True)
    else:
        win = build_pairwise_sql(metrics, "A", "B", "win")
        loss = build_pairwise_sql(metrics, "A", "B", "loss")
        # SP = how many seasons (rows) this player has in total, including this one -
        # every one of those is excluded from the comparison pool, not just this exact row.
        q = f"""
            SELECT A.Player, A.{t_col} as Year,
                   (SELECT COUNT(*) FROM {disc}) as TR,
                   (SELECT COUNT(*) FROM {disc} B WHERE B.Player = A.Player) as SP,
                   (SELECT COUNT(*) FROM {disc} B WHERE B.Player != A.Player AND {win} >= 2) as WC,
                   (SELECT COUNT(*) FROM {disc} B WHERE B.Player != A.Player AND {loss} >= 2) as LC,
                   (SELECT COUNT(*) FROM {disc} B WHERE B.Player != A.Player AND {loss} = 3) as LC3
            FROM {disc} A
        """
        df = pd.read_sql(q, conn)
        df['Others'] = df['TR'] - df['SP']  # comparison pool: every OTHER player's seasons
        df['Ties'] = df['Others'] - df['WC'] - df['LC']
        df['Wins'] = df['WC']
        df['Win %'] = df.apply(lambda r: pct(r['WC'], r['Others']), axis=1)
        df['Losses'] = df['LC']
        df['Loss %'] = df.apply(lambda r: pct(r['LC'], r['Others']), axis=1)
        df['3-0 Losses'] = df['LC3']
        df['Tie %'] = df.apply(lambda r: pct(r['Ties'], r['Others']), axis=1)
        display_cols = ['Player', 'Year', 'Wins', 'Win %', 'Losses', 'Loss %', '3-0 Losses', 'Ties', 'Tie %']
        st.dataframe(df.sort_values("Wins", ascending=False)[display_cols].reset_index(drop=True), use_container_width=True, hide_index=True)

# --- TAB 4: DETAILS ---
elif st.session_state.nav_choice == "👤 Player Details":
    all_p = sorted(list(set(pd.read_sql("SELECT Player FROM batting", conn)['Player']) | set(pd.read_sql("SELECT Player FROM bowling", conn)['Player'])))
    target = st.selectbox("Select Player", all_p)
    if target:
        for t, lab in [("batting", "Batting"), ("bowling", "Bowling")]:
            metrics = st.session_state.bat_metrics if t == "batting" else st.session_state.bowl_metrics
            cols_sql = ", ".join(q_col(m['col']) for m in metrics)
            win = build_pairwise_sql(metrics, "A", "B", "win")
            loss = build_pairwise_sql(metrics, "A", "B", "loss")
            q = f"SELECT A.{t_col} as Year, {cols_sql}, (SELECT COUNT(*) FROM {t}) as TR, (SELECT COUNT(*) FROM {t} B WHERE {win} >= 2) as WC, (SELECT COUNT(*) FROM {t} B WHERE {loss} >= 2) as LC FROM {t} A WHERE A.Player = '{target}'"
            df = pd.read_sql(q, conn)
            if not df.empty:
                st.subheader(lab)
                df['Wins'] = df['WC']
                df['Win %'] = df.apply(lambda r: pct(r['WC'], r['TR'] - 1), axis=1)
                df['Losses'] = df['LC']
                df['Loss %'] = df.apply(lambda r: pct(r['LC'], r['TR'] - 1), axis=1)
                df['Ties'] = df['TR'] - df['WC'] - df['LC'] - 1
                df['Tie %'] = df.apply(lambda r: pct(r['Ties'], r['TR'] - 1), axis=1)
                evt = st.dataframe(df.drop(columns=['TR', 'WC', 'LC']), use_container_width=True, hide_index=True, on_select="rerun", selection_mode="single-row", key=f"det_scr_{t}")
                if evt.selection.rows:
                    sel = df.iloc[evt.selection.rows[0]]
                    thresholds_key = f"{t}_thresholds"
                    st.session_state.setdefault(thresholds_key, {})
                    for m in metrics:
                        st.session_state[thresholds_key][m['col']] = sel[m['col']]
                    st.session_state.nav_choice = "Batting Milestones" if t == "batting" else "Bowling Milestones"
                    st.rerun()

# --- TAB 5: SQUAD ---
elif st.session_state.nav_choice == "🏟️ Squad Comparison":
    with st.expander("🛠️ Manage Squads"):
        c1, c2 = st.columns(2)
        if c1.button("🗑️ Clear Squads"): st.session_state.squad_a, st.session_state.squad_b = [], []; st.rerun()
        c2.text_input("📋 Squad Code", value=json.dumps({"a": st.session_state.squad_a, "b": st.session_state.squad_b}))
        load = st.text_input("📥 Paste Code:")
        if st.button("🔄 Execute Load") and load:
            try: d = json.loads(load); st.session_state.squad_a, st.session_state.squad_b = d.get('a', []), d.get('b', []); st.rerun()
            except Exception: st.error("ERR")
    all_p = sorted(list(set(pd.read_sql("SELECT Player FROM batting", conn)['Player']) | set(pd.read_sql("SELECT Player FROM bowling", conn)['Player'])))
    c1, c2 = st.columns(2)
    with c1:
        n = st.selectbox("Add to A", [""] + all_p, key="sqa"); (st.session_state.squad_a.append(n), st.rerun()) if n and n not in st.session_state.squad_a else None
        st.session_state.squad_a = st.multiselect("Squad A", st.session_state.squad_a, default=st.session_state.squad_a)
    with c2:
        n = st.selectbox("Add to B", [""] + all_p, key="sqb"); (st.session_state.squad_b.append(n), st.rerun()) if n and n not in st.session_state.squad_b else None
        st.session_state.squad_b = st.multiselect("Squad B", st.session_state.squad_b, default=st.session_state.squad_b)

    with st.expander("🚫 Omit Specific Player-Years"):
        st.caption(
            "Exclude individual player-years from a squad's eligible pool without dropping the player "
            "entirely - e.g. leave Amla's 2010 batting year out of Squad A while keeping his other years."
        )
        for squad_label, squad_list, omit_key in [
            ("Squad A", st.session_state.squad_a, "squad_a_omit"),
            ("Squad B", st.session_state.squad_b, "squad_b_omit"),
        ]:
            st.markdown(f"**{squad_label}**")
            if not squad_list:
                st.caption("No players in this squad yet.")
                continue
            option_map = {}
            for table, disc_label in [("batting", "Batting"), ("bowling", "Bowling")]:
                l = "('" + "','".join(squad_list) + "')"
                py_df = pd.read_sql(f"SELECT Player, {t_col} as Year FROM {table} WHERE Player IN {l}", conn)
                for _, r in py_df.iterrows():
                    label = f"{r['Player']} — {r['Year']} ({disc_label})"
                    option_map[label] = (table, r['Player'], r['Year'])
            current_labels = [lbl for lbl, val in option_map.items() if val in st.session_state[omit_key]]
            chosen = st.multiselect(
                f"Years to omit from {squad_label}", sorted(option_map.keys()),
                default=current_labels, key=f"{omit_key}_select",
            )
            st.session_state[omit_key] = {option_map[c] for c in chosen}

    if st.session_state.squad_a and st.session_state.squad_b:
        sub = st.radio("Mode:", ["Individual Benchmark", "Squad Pairwise"], horizontal=True)
        if sub == "Individual Benchmark":
            d_dir = st.radio("Direction:", ["Squad A ➡️ B", "Squad B ➡️ A"], horizontal=True)
            src, trg = (st.session_state.squad_a, st.session_state.squad_b) if "A ➡️" in d_dir else (st.session_state.squad_b, st.session_state.squad_a)
            src_omit, trg_omit = (st.session_state.squad_a_omit, st.session_state.squad_b_omit) if "A ➡️" in d_dir else (st.session_state.squad_b_omit, st.session_state.squad_a_omit)
            p = st.selectbox("Pick Benchmark Player:", src, key=f"sq_p_sel_{d_dir}_{len(src)}")
            if p:
                b_y_all = pd.read_sql(f"SELECT {t_col} as Year FROM batting WHERE Player='{p}'", conn)['Year'].tolist()
                w_y_all = pd.read_sql(f"SELECT {t_col} as Year FROM bowling WHERE Player='{p}'", conn)['Year'].tolist()
                b_y = [y for y in b_y_all if ("batting", p, y) not in src_omit]
                w_y = [y for y in w_y_all if ("bowling", p, y) not in src_omit]
                disc = st.radio("Type:", (["Batting"] if b_y else []) + (["Bowling"] if w_y else []), horizontal=True)
                y = st.selectbox("Year:", b_y if disc == "Batting" else w_y)
                if y:
                    table = "batting" if disc == "Batting" else "bowling"
                    metrics = st.session_state.bat_metrics if disc == "Batting" else st.session_state.bowl_metrics
                    bench = pd.read_sql(f"SELECT * FROM {table} WHERE Player='{p}' AND {t_col}='{y}'", conn).iloc[0]
                    st.info(f"📍 Benchmark: {p} ({y})")
                    met_cols = st.columns(len(metrics))
                    thresholds = []
                    for i, m in enumerate(metrics):
                        val = bench[m['col']]
                        met_cols[i].metric(m['col'], val)
                        thresholds.append(val)
                    target_str = "('" + "','".join(trg) + "')"
                    target_omit_sql = build_omit_sql(t_col, trg_omit, table)
                    select_cols = ", ".join(q_col(m['col']) for m in metrics)
                    wins_a = build_case_sql(metrics, thresholds, 'win')
                    ties_a = build_case_sql(metrics, thresholds, 'tie')
                    losses_a = build_case_sql(metrics, thresholds, 'loss')
                    q = f"""
                        SELECT Player, {t_col} as Year, {select_cols},
                               ({wins_a}) as WinsA, ({ties_a}) as TiesA, ({losses_a}) as LossesA
                        FROM {table} WHERE Player IN {target_str} {target_omit_sql}
                        ORDER BY WinsA DESC
                    """
                    display_styled_results(pd.read_sql(q, conn), f"Against {p}")
        else:
            t_disc = st.radio("Choose Discipline:", ["Batting", "Bowling"], horizontal=True).lower()
            metrics = st.session_state.bat_metrics if t_disc == "batting" else st.session_state.bowl_metrics
            a_l = "('" + "','".join(st.session_state.squad_a) + "')"
            b_l = "('" + "','".join(st.session_state.squad_b) + "')"
            omit_a_outer = build_omit_sql(t_col, st.session_state.squad_a_omit, t_disc, alias="A")
            omit_b_outer = build_omit_sql(t_col, st.session_state.squad_b_omit, t_disc, alias="A")
            omit_a_inner = build_omit_sql(t_col, st.session_state.squad_a_omit, t_disc, alias="B")
            omit_b_inner = build_omit_sql(t_col, st.session_state.squad_b_omit, t_disc, alias="B")
            win = build_pairwise_sql(metrics, "A", "B", "win")
            loss = build_pairwise_sql(metrics, "A", "B", "loss")
            c1, c2 = st.columns(2)
            with c1:
                st.write("Squad A vs B")
                q_a = f"SELECT A.Player, A.{t_col} as Year, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {b_l} {omit_b_inner}) as TR, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {b_l} {omit_b_inner} AND {win} >= 2) as WC, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {b_l} {omit_b_inner} AND {loss} >= 2) as LC FROM {t_disc} A WHERE A.Player IN {a_l} {omit_a_outer}"
                df_a = pd.read_sql(q_a, conn)
                df_a['Wins'] = df_a['WC']
                df_a['Win %'] = df_a.apply(lambda r: pct(r['WC'], r['TR']), axis=1)
                df_a['Losses'] = df_a['LC']
                df_a['Loss %'] = df_a.apply(lambda r: pct(r['LC'], r['TR']), axis=1)
                df_a['Ties'] = df_a['TR'] - df_a['WC'] - df_a['LC']
                df_a['Tie %'] = df_a.apply(lambda r: pct(r['Ties'], r['TR']), axis=1)
                st.dataframe(df_a[['Player', 'Year', 'Wins', 'Win %', 'Losses', 'Loss %', 'Ties', 'Tie %']], hide_index=True)
            with c2:
                st.write("Squad B vs A")
                q_b = f"SELECT A.Player, A.{t_col} as Year, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {a_l} {omit_a_inner}) as TR, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {a_l} {omit_a_inner} AND {win} >= 2) as WC, (SELECT COUNT(*) FROM {t_disc} B WHERE B.Player IN {a_l} {omit_a_inner} AND {loss} >= 2) as LC FROM {t_disc} A WHERE A.Player IN {b_l} {omit_b_outer}"
                df_b = pd.read_sql(q_b, conn)
                df_b['Wins'] = df_b['WC']
                df_b['Win %'] = df_b.apply(lambda r: pct(r['WC'], r['TR']), axis=1)
                df_b['Losses'] = df_b['LC']
                df_b['Loss %'] = df_b.apply(lambda r: pct(r['LC'], r['TR']), axis=1)
                df_b['Ties'] = df_b['TR'] - df_b['WC'] - df_b['LC']
                df_b['Tie %'] = df_b.apply(lambda r: pct(r['Ties'], r['TR']), axis=1)
                st.dataframe(df_b[['Player', 'Year', 'Wins', 'Win %', 'Losses', 'Loss %', 'Ties', 'Tie %']], hide_index=True)

# --- TAB 6: FORMAT ANALYSIS ---
elif st.session_state.nav_choice == "🧬 Format Analysis":
    sub6 = st.radio("Feature:", ["🛡️ Unbeatable", "🔍 Group Killers"], horizontal=True)
    t = st.radio("Disc:", ["Batting", "Bowling"], horizontal=True, key="disc_fa").lower()
    metrics = st.session_state.bat_metrics if t == "batting" else st.session_state.bowl_metrics
    pl = sorted(pd.read_sql(f"SELECT DISTINCT Player FROM {t}", conn)['Player'].tolist())
    if "Unbeatable" in sub6:
        cx1, cx2 = st.columns(2)
        with cx1: o1 = st.multiselect("Omit Pool:", pl, key="om1")
        with cx2: o2 = st.multiselect("Omit Check:", pl, key="om2")
        k_v = st.radio("Size (K):", [1, 2, 3], horizontal=True)
        if st.button("🚀 Find"):
            with st.spinner("Analyzing..."):
                df_f = pd.read_sql(f"SELECT * FROM {t}", conn); df_k = df_f[~df_f['Player'].isin(o2)]; cand = [p for p in pl if p not in o1]
                if not cand: st.error("No players.")
                else:
                    kill_sets = []
                    for _, rx in df_k.iterrows():
                        beaten = set()
                        for pn in cand:
                            py = df_f[df_f['Player'] == pn]; b = 0
                            for _, rp in py.iterrows():
                                if compare_rows(rx, rp, metrics) >= 2: b += 1
                            if b == len(py): beaten.add(pn)
                        kill_sets.append(beaten)
                    res = []
                    for combo in combinations(cand, k_v):
                        cs = set(combo); unb = True
                        for ks in kill_sets:
                            if cs.issubset(ks): unb = False; break
                        if unb: res.append(list(combo))
                    if res: st.dataframe(pd.DataFrame(res, columns=[f"P{i+1}" for i in range(k_v)]), hide_index=True)
                    else: st.error("None found.")
    else:
        target = st.multiselect("Select Group:", pl)
        if st.button("🔎 Find Killers"):
            if target:
                df_full = pd.read_sql(f"SELECT * FROM {t}", conn); kl = []
                for _, rx in df_full.iterrows():
                    e = True
                    for pn in target:
                        py = df_full[df_full['Player'] == pn]; pb = True
                        for _, rp in py.iterrows():
                            if compare_rows(rx, rp, metrics) < 2: pb = False; break
                        if not pb: e = False; break
                    if e: kl.append(rx)
                if kl: st.dataframe(pd.DataFrame(kl).rename(columns={t_col: 'Year'}).reset_index(drop=True), hide_index=True)
                else: st.error("No killers.")

# --- TAB 7: EDIT DATA ---
elif st.session_state.nav_choice == "✏️ Edit Data":
    st.caption(
        f"Editing '{st.session_state.active_dataset}' directly. Changes save to that dataset's file immediately — "
        "if you want them to survive a restart/redeploy, use '📌 Make Permanent' in the sidebar afterwards."
    )
    edit_table = st.radio("Table:", ["Batting", "Bowling"], horizontal=True, key="edit_table_choice").lower()

    # Edits always target the RAW file connection, never the year/country-filtered
    # in-memory one — otherwise saves would vanish on the next rerun.
    full_df = pd.read_sql(f"SELECT * FROM {edit_table}", raw_conn)
    search = st.text_input("Search player (optional — narrows what's shown/edited, doesn't affect the rest of the data)", key="edit_search")
    display_df = full_df[full_df['Player'].str.contains(search, case=False, na=False)] if search else full_df
    st.caption(f"Showing {len(display_df)} of {len(full_df)} total rows.")

    edited_df = st.data_editor(
        display_df, use_container_width=True, hide_index=True,
        num_rows="fixed",  # editing existing values only — add/remove rows not yet supported here
        key=f"data_editor_{edit_table}_{st.session_state.active_dataset}",
    )

    if st.button("💾 Save Changes", key=f"save_edit_{edit_table}"):
        full_df.loc[edited_df.index] = edited_df
        full_df.to_sql(edit_table, raw_conn, index=False, if_exists='replace')
        st.success(f"Saved changes to '{edit_table}' in '{st.session_state.active_dataset}'.")
        st.rerun()

# --- TAB 8: AUCTION SHEET ---
elif st.session_state.nav_choice == "🏆 Auction Sheet":
    st.caption(
        "Upload an auction-order file (PDF or TXT) and the Global Rankings CSVs (Player, Year, Wins %, Losses, Ties) "
        "for whichever formats apply to this tournament. Output is ordered exactly like your auction list, "
        "with separate batting/bowling loss-threshold counts and a best-year rank per format."
    )
    st.caption(
        "💡 If the PDF isn't picking up players correctly (e.g. it shows page footers instead of names), "
        "paste the same list into a plain .txt file (one player per line, numbered or not) and upload that instead — "
        "it's a lot more reliable than PDF text extraction."
    )

    THRESHOLDS = [5, 10, 20, 30, 40, 50]
    pdf_file = st.file_uploader("Auction Order (PDF or TXT)", type=["pdf", "txt"], key="auction_pdf")

    formats_config = {}
    for fmt_name in ["Test", "ODI", "T20I"]:
        include = st.checkbox(f"Include {fmt_name}", key=f"include_{fmt_name}")
        if include:
            c1, c2 = st.columns(2)
            bat_file = c1.file_uploader(f"{fmt_name} Batting CSV", type=["csv"], key=f"{fmt_name}_bat_csv")
            bowl_file = c2.file_uploader(f"{fmt_name} Bowling CSV", type=["csv"], key=f"{fmt_name}_bowl_csv")
            formats_config[fmt_name] = {"bat": bat_file, "bowl": bowl_file}

    if st.button("🚀 Generate Auction Sheet"):
        if not pdf_file:
            st.error("Upload the auction order file first.")
        elif not formats_config:
            st.error("Include at least one format.")
        else:
            missing = [f for f, cfg in formats_config.items() if cfg['bat'] is None or cfg['bowl'] is None]
            if missing:
                st.error(f"Missing a batting or bowling CSV for: {', '.join(missing)}")
            else:
                with st.spinner("Matching players and building the sheet..."):
                    auction_players = parse_auction_order_bytes(pdf_file.read(), pdf_file.name)
                    loaded = _load_format_dataframes(formats_config)
                    indices = _build_indices(loaded)
                    auto_records = _auto_match_records(auction_players, indices)
                # Stashed in session_state so the manual-correction controls below
                # can rebuild the sheet on every tweak without needing the original
                # file uploads to still be present, and survive the reruns that
                # each correction dropdown triggers.
                st.session_state.auction_players = auction_players
                st.session_state.auction_loaded = loaded
                st.session_state.auction_indices = indices
                st.session_state.auction_auto_records = auto_records
                st.session_state.auction_format_names = list(formats_config.keys())
                st.session_state.auction_thresholds = THRESHOLDS
                st.session_state.auction_overrides = {}  # fresh generate clears prior corrections

    if st.session_state.get("auction_players"):
        auction_players = st.session_state.auction_players
        loaded = st.session_state.auction_loaded
        indices = st.session_state.auction_indices
        auto_records = st.session_state.auction_auto_records
        format_names = st.session_state.auction_format_names
        thresholds = st.session_state.auction_thresholds
        overrides = st.session_state.setdefault("auction_overrides", {})

        # Cells worth showing a corrector for: ambiguous/weak matches that
        # cross-sheet propagation couldn't resolve either, plus any "no match
        # found" case where this player wasn't identified in ANY
        # format/discipline (see _needs_review for why a no-match
        # elsewhere-confirmed player is skipped).
        resolved, confirmed_counts = _resolve_all(auction_players, format_names, indices, auto_records, overrides)
        flagged_keys = sorted(
            key for key, (matched, status) in resolved.items()
            if _needs_review(status, key[0], confirmed_counts)
        )

        if flagged_keys:
            with st.expander(f"⚠️ {len(flagged_keys)} match(es) need review", expanded=True):
                st.caption(
                    "Pick the correct player for any flagged cell below. The sheet and Excel file "
                    "below update immediately - no extra step needed."
                )
                for order, fmt_name, disc in flagged_keys:
                    rec = auto_records[(order, fmt_name, disc)]
                    disc_label = "Bat" if disc == "bat" else "Bowl"
                    df = loaded[(fmt_name, disc)]
                    all_names = sorted(df['Player'].dropna().unique().tolist())
                    options = [AUCTION_KEEP_AUTO] + all_names + [AUCTION_EXCLUDE]

                    current_override = overrides.get((order, fmt_name, disc))
                    if current_override is None:
                        default_label = AUCTION_KEEP_AUTO
                    else:
                        default_label = current_override
                    default_index = options.index(default_label) if default_label in options else 0

                    auto_desc = (
                        f"auto-matched to '{rec['matched_name']}' ({rec['status']}, please verify)"
                        if rec['matched_name'] else "no automatic match found"
                    )
                    st.markdown(f"**#{order} {rec['auction_player']} — {fmt_name} {disc_label}** _({auto_desc})_")
                    choice = st.selectbox(
                        "Correct player", options, index=default_index,
                        key=f"auction_override_{order}_{fmt_name}_{disc}",
                        label_visibility="collapsed",
                    )
                    if choice == AUCTION_KEEP_AUTO:
                        overrides.pop((order, fmt_name, disc), None)
                    else:
                        overrides[(order, fmt_name, disc)] = choice
                    st.divider()

                if overrides and st.button("↩️ Reset All Corrections"):
                    st.session_state.auction_overrides = {}
                    st.rerun()

        result_df, warnings = build_auction_sheet_df(auction_players, format_names, loaded, resolved, confirmed_counts, thresholds)
        st.success(f"Built the sheet for {len(result_df)} players" + (f" ({len(overrides)} manually corrected)." if overrides else "."))

        locked_cols = ["Auction Order", "Auction Player"]
        optional_cols = [c for c in result_df.columns if c not in locked_cols]
        selected_cols = st.multiselect(
            "Columns to include (Auction Order & Auction Player are always included)",
            optional_cols, default=optional_cols, key="auction_col_select",
        )
        output_df = result_df[locked_cols + [c for c in optional_cols if c in selected_cols]]

        buf = io.BytesIO()
        with pd.ExcelWriter(buf, engine="openpyxl") as writer:
            output_df.to_excel(writer, index=False, sheet_name="Auction Sheet")
            ws = writer.sheets["Auction Sheet"]
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            header_font = Font(name="Arial", bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
            body_font = Font(name="Arial", size=10)
            for col_idx, col_name in enumerate(output_df.columns, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.column_dimensions[get_column_letter(col_idx)].width = min(max(len(str(col_name)) + 2, 10), 22)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.font = body_font
            ws.freeze_panes = "C2"
            ws.row_dimensions[1].height = 30

        st.download_button(
            "⬇️ Download Auction Sheet (.xlsx)", data=buf.getvalue(),
            file_name="auction_sheet.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key="auction_download",
        )
        st.dataframe(output_df, use_container_width=True, hide_index=True)

conn.close()
if raw_conn is not conn:
    raw_conn.close()
